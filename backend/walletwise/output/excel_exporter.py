from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from walletwise.models import ExtractionResult, Transaction


class ExcelExporter:
    """Bank-independent WalletWise Excel exporter."""

    SCHEMA_VERSION = "1.3"
    HEADERS = [
        "Date",
        "Name",
        "Original Description",
        "Category",
        "Category Confidence",
        "Category Reason",
        "Financial Role",
        "Role Confidence",
        "Role Reason",
        "Amount",
        "Type",
        "Balance",
        "Bank Charges",
        "Bank",
        "Extraction Confidence",
        "Name Confidence",
        "Name Reason",
        "Review Notes",
    ]

    def export(self, result: ExtractionResult, output_file: str):
        workbook = Workbook()
        workbook.remove(workbook.active)

        header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        green_fill = PatternFill(fill_type="solid", fgColor="C6EFCE")
        orange_fill = PatternFill(fill_type="solid", fgColor="FCE4D6")
        red_fill = PatternFill(fill_type="solid", fgColor="FFC7CE")

        transactions_by_month = defaultdict(list)
        for transaction in result.transactions:
            month_key = transaction.date.strftime("%Y-%m")
            transactions_by_month[month_key].append(transaction)

        confidence_counts = {"high": 0, "medium": 0, "low": 0}
        naming_counts = {"high": 0, "medium": 0, "low": 0, "unknown": 0}

        for month_key in sorted(transactions_by_month):
            month_date = datetime.strptime(month_key, "%Y-%m")
            worksheet = workbook.create_sheet(month_date.strftime("%B %Y"))
            worksheet.append(self.HEADERS)

            for cell in worksheet[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")

            for transaction in transactions_by_month[month_key]:
                worksheet.append(self._row(transaction))
                excel_row = worksheet.max_row

                # -------------------------
                # Whole-row extraction confidence
                # -------------------------
                extraction_level = transaction.confidence
                confidence_counts[extraction_level] = confidence_counts.get(extraction_level, 0) + 1

                row_fill = {
                    "high": green_fill,
                    "medium": orange_fill,
                    "low": red_fill,
                }.get(extraction_level, orange_fill)

                for cell in worksheet[excel_row]:
                    cell.fill = row_fill

                # -------------------------
                # Name confidence overrides ONLY the Name cell.
                # Green = known/strong name
                # Orange = usable but needs confirmation
                # Red = no usable name
                # -------------------------
                naming_level = transaction.naming_confidence or "unknown"
                naming_counts[naming_level] = naming_counts.get(naming_level, 0) + 1

                name_fill = {
                    "high": green_fill,
                    "medium": orange_fill,
                    "low": orange_fill,
                    "unknown": orange_fill,
                }.get(naming_level, orange_fill)

                if not (transaction.display_name or transaction.raw_description):
                    name_fill = red_fill

                worksheet.cell(excel_row, 2).fill = name_fill

                category_fill = {
                    "high": green_fill,
                    "medium": orange_fill,
                    "low": orange_fill,
                    "unknown": orange_fill,
                }.get(transaction.category_confidence or "unknown", orange_fill)
                if transaction.category == "Uncategorised":
                    category_fill = orange_fill
                worksheet.cell(excel_row, 4).fill = category_fill

                worksheet.cell(excel_row, 1).number_format = "dd mmm yyyy"
                role_fill = {
                    'high': green_fill,
                    'medium': orange_fill,
                    'low': orange_fill,
                    'unknown': orange_fill,
                }.get(transaction.role_confidence or 'unknown', orange_fill)
                worksheet.cell(excel_row, 7).fill = role_fill

                worksheet.cell(excel_row, 10).number_format = 'R #,##0.00;[Red]-R #,##0.00'
                worksheet.cell(excel_row, 12).number_format = 'R #,##0.00;[Red]-R #,##0.00'
                worksheet.cell(excel_row, 13).number_format = 'R #,##0.00;[Red]-R #,##0.00'

            widths = {
                "A": 15, "B": 30, "C": 65, "D": 22, "E": 20, "F": 55,
                "G": 18, "H": 18, "I": 55, "J": 16, "K": 18, "L": 16,
                "M": 18, "N": 14, "O": 22, "P": 18, "Q": 65, "R": 60,
            }
            for column, width in widths.items():
                worksheet.column_dimensions[column].width = width

            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = f"A1:R{worksheet.max_row}"

        summary = workbook.create_sheet("Extraction Summary")
        summary.append(["Metric", "Value"])
        summary["A1"].font = Font(bold=True)
        summary["B1"].font = Font(bold=True)

        summary.append(["WalletWise schema version", self.SCHEMA_VERSION])
        summary.append(["Detected bank", result.bank])
        summary.append(["Transactions extracted", len(result.transactions)])
        summary.append(["Extraction high confidence (green)", confidence_counts.get("high", 0)])
        summary.append(["Extraction medium confidence (orange)", confidence_counts.get("medium", 0)])
        summary.append(["Extraction low confidence (red)", confidence_counts.get("low", 0)])
        summary.append(["Name high confidence (green)", naming_counts.get("high", 0)])
        summary.append([
            "Names needing confirmation (orange)",
            naming_counts.get("medium", 0)
            + naming_counts.get("low", 0)
            + naming_counts.get("unknown", 0),
        ])
        summary.append(["Unassigned rows", len(result.unassigned_rows)])
        summary.append([])
        summary.append(["Page", "Visual rows detected"])

        for page_number in sorted(result.page_counts):
            summary.append([page_number, result.page_counts[page_number]])

        if result.debug_files:
            summary.append([])
            summary.append(["Debug files"])
            for filename in result.debug_files:
                summary.append([filename])

        if result.unassigned_rows:
            summary.append([])
            summary.append([
                "Unassigned Bank",
                "Page",
                "Raw Date",
                "Description",
                "Amount",
                "Balance",
                "Reason",
            ])

            for row in result.unassigned_rows:
                summary.append([
                    row.bank,
                    row.page,
                    row.raw_date,
                    row.raw_description,
                    row.raw_amount,
                    row.raw_balance,
                    row.reason,
                ])

        summary.column_dimensions["A"].width = 40
        summary.column_dimensions["B"].width = 24
        summary.column_dimensions["C"].width = 18
        summary.column_dimensions["D"].width = 60
        summary.column_dimensions["E"].width = 18
        summary.column_dimensions["F"].width = 18
        summary.column_dimensions["G"].width = 60

        workbook.save(output_file)
        return confidence_counts

    @staticmethod
    def _row(transaction: Transaction):
        date_value = None if transaction.transaction_type == "closing_balance" else transaction.date

        return [
            date_value,
            transaction.display_name or transaction.raw_description,
            transaction.raw_description,
            transaction.category,
            transaction.category_confidence.title(),
            transaction.category_reason,
            transaction.financial_role.title(),
            transaction.role_confidence.title(),
            transaction.role_reason,
            transaction.amount,
            transaction.transaction_type,
            transaction.balance,
            transaction.bank_charges,
            transaction.bank,
            transaction.confidence.title(),
            transaction.naming_confidence.title(),
            transaction.naming_reason,
            "; ".join(transaction.review_reasons),
        ]
